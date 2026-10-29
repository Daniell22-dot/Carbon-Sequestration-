"""Word (.docx) report generation with python-docx."""

from __future__ import annotations

from collections.abc import Mapping, Sequence
from pathlib import Path

import pandas as pd

from src.utils.io import ensure_dir
from src.utils.logging import get_logger

logger = get_logger(__name__)

FigureSpec = Path | str | Sequence[Path | str]
TableSpec = pd.DataFrame | Sequence[pd.DataFrame]


def _add_table(document: object, frame: pd.DataFrame, max_rows: int = 40) -> None:
    from docx.shared import Pt

    subset = frame.head(max_rows)
    table = document.add_table(rows=1, cols=len(subset.columns))  # type: ignore[attr-defined]
    table.style = "Light Grid Accent 1"
    header_cells = table.rows[0].cells
    for position, column in enumerate(subset.columns):
        header_cells[position].text = str(column)
        for paragraph in header_cells[position].paragraphs:
            for run in paragraph.runs:
                run.bold = True
                run.font.size = Pt(9)

    for _, row in subset.iterrows():
        cells = table.add_row().cells
        for position, column in enumerate(subset.columns):
            value = row[column]
            text = f"{value:.4g}" if isinstance(value, float) else str(value)
            cells[position].text = text
            for paragraph in cells[position].paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(9)


def create_report(
    title: str,
    sections: Mapping[str, str],
    figures: Mapping[str, FigureSpec] | None = None,
    tables: Mapping[str, TableSpec] | None = None,
    output_path: Path | str = "output/reports/report.docx",
    subtitle: str | None = None,
    metadata: Mapping[str, str] | None = None,
) -> Path:
    """Generate a professional Word report.

    Args:
        title: Report title.
        sections: Mapping of section heading to body text. Body text is split
            on blank lines into paragraphs.
        figures: Optional mapping of section heading to one or more figure
            paths. A figure is inserted after that section's body.
        tables: Optional mapping of section heading to one or more DataFrames.
        output_path: Destination ``.docx`` path.
        subtitle: Optional subtitle shown on the title page.
        metadata: Optional provenance lines shown on the title page.

    Returns:
        The destination path.

    Raises:
        ImportError: If python-docx is not installed.
        FileNotFoundError: If a referenced figure does not exist.
    """
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Inches, Pt
    except ImportError as error:  # pragma: no cover
        raise ImportError("python-docx is required to generate Word reports") from error

    document = Document()

    title_paragraph = document.add_heading(title, level=0)
    title_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if subtitle:
        subtitle_paragraph = document.add_paragraph(subtitle)
        subtitle_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in subtitle_paragraph.runs:
            run.font.size = Pt(13)
            run.italic = True

    if metadata:
        document.add_paragraph()
        for key, value in metadata.items():
            line = document.add_paragraph(f"{key}: {value}")
            line.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in line.runs:
                run.font.size = Pt(9)

    document.add_page_break()
    document.add_heading("Table of Contents", level=1)
    for section_name in sections:
        document.add_paragraph(section_name, style="List Bullet")
    document.add_page_break()

    for section_name, content in sections.items():
        document.add_heading(section_name, level=1)
        for paragraph_text in str(content).split("\n\n"):
            if paragraph_text.strip():
                document.add_paragraph(paragraph_text.strip())

        section_figures = figures.get(section_name) if figures else None
        if section_figures:
            specs = (
                [section_figures]
                if isinstance(section_figures, (str, Path))
                else list(section_figures)
            )
            for spec in specs:
                figure_path = Path(spec)
                if not figure_path.exists():
                    raise FileNotFoundError(f"Figure not found: {figure_path}")
                document.add_picture(str(figure_path), width=Inches(6.0))
                document.add_paragraph(f"Figure: {figure_path.stem}").alignment = (
                    WD_ALIGN_PARAGRAPH.CENTER
                )

        section_tables = tables.get(section_name) if tables else None
        if section_tables is not None:
            specs_tables = (
                [section_tables]
                if isinstance(section_tables, pd.DataFrame)
                else list(section_tables)
            )
            for frame in specs_tables:
                _add_table(document, frame)
                document.add_paragraph()

    destination = Path(output_path)
    ensure_dir(destination.parent)
    document.save(str(destination))
    logger.info("Word report written to {}", destination)
    return destination
